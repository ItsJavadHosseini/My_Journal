
"""Import MetaTrader Excel reports into the Trading Journal domain model."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Sequence
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from trading_journal.domain.models import Trade, TradeDirection
from trading_journal.imports.base import (
    ImportResult,
    ImportValidationError,
    TradeImporter,
)


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

SECTION_TITLE = "positions"

SECTION_TERMINATORS = {
    "orders",
    "deals",
    "results",
    "history",
    "open positions",
    "closed positions",
}

FOOTER_PREFIXES = (
    "balance:",
    "equity:",
    "credit facility:",
    "margin:",
    "free margin:",
)

HEADER_ALIASES = {
    "position": {"position", "positionid", "ticket"},
    "symbol": {"symbol", "instrument"},
    "type": {"type", "direction", "side"},
    "volume": {"volume", "lots", "size"},
    "sl": {"sl", "s/l", "stop loss"},
    "tp": {"tp", "t/p", "take profit"},
    "commission": {"commission", "comm"},
    "swap": {"swap", "swaps"},
    "profit": {"profit", "p/l", "pnl"},
}

DATETIME_FORMATS = (
    "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y.%m.%d",
    "%Y-%m-%d",
)

SUPPORTED_EXTENSIONS = {".xlsx"}


# ---------------------------------------------------------------------------
# Internal data structures
# ---------------------------------------------------------------------------

@dataclass(frozen=True, slots=True)
class _Section:
    worksheet: Worksheet
    title_row: int
    header_row: int
    data_start_row: int
    data_end_row: int  # Exclusive


@dataclass(frozen=True, slots=True)
class _Columns:
    entry_time: int
    position_id: int
    symbol: int
    direction: int
    volume: int
    entry_price: int
    exit_time: int
    exit_price: int
    stop_loss: int | None
    take_profit: int | None
    commission: int | None
    swap: int | None
    profit: int


# ---------------------------------------------------------------------------
# Normalization helpers
# ---------------------------------------------------------------------------

def _text(value: Any) -> str:
    """Convert a cell value to a trimmed string."""
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def _normalize(value: Any) -> str:
    """Normalize a header or section title for comparison."""
    return re.sub(r"\s+", " ", _text(value).lower()).strip()


def _normalize_header(value: Any) -> str:
    """Normalize header names while ignoring punctuation and spaces."""
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def _is_blank_row(row: Sequence[Any]) -> bool:
    return all(value is None or _text(value) == "" for value in row)


def _is_footer_row(row: Sequence[Any]) -> bool:
    """Identify common report summary rows."""
    for value in row:
        cell = _normalize(value)
        if any(cell.startswith(prefix) for prefix in FOOTER_PREFIXES):
            return True
    return False


def _section_name(row: Sequence[Any]) -> str | None:
    """Return a recognized section title if the row contains one."""
    for value in row:
        normalized = _normalize(value)
        if normalized in SECTION_TERMINATORS:
            return normalized
    return None


# ---------------------------------------------------------------------------
# Excel section and header detection
# ---------------------------------------------------------------------------

def _find_positions_section(worksheet: Worksheet) -> _Section:
    """
    Find the Positions section and its header without fixed cell addresses.

    A valid header must contain the expected core columns and two Time and
    Price columns, as in the standard MetaTrader Positions report.
    """
    rows = worksheet.iter_rows(values_only=True)
    materialized_rows = [tuple(row) for row in rows]

    title_indices = [
        index
        for index, row in enumerate(materialized_rows)
        if any(_normalize(cell) == SECTION_TITLE for cell in row)
    ]

    if not title_indices:
        raise ImportValidationError(
            f"Section 'Positions' was not found in worksheet "
            f"{worksheet.title!r}."
        )

    if len(title_indices) > 1:
        raise ImportValidationError(
            f"Multiple 'Positions' section titles were found in worksheet "
            f"{worksheet.title!r}; refusing to guess which one to import."
        )

    title_index = title_indices[0]
    header_index: int | None = None

    for index in range(title_index + 1, len(materialized_rows)):
        row = materialized_rows[index]

        if _is_blank_row(row):
            continue

        if _section_name(row) is not None:
            break

        if _is_positions_header(row):
            header_index = index
            break

    if header_index is None:
        raise ImportValidationError(
            f"Could not find a valid header row after 'Positions' in "
            f"worksheet {worksheet.title!r}. Expected Position, Symbol, "
            f"Type, Volume, two Time columns, two Price columns, and Profit."
        )

    end_index = len(materialized_rows)
    consecutive_blank_rows = 0

    for index in range(header_index + 1, len(materialized_rows)):
        row = materialized_rows[index]

        if _section_name(row) is not None or _is_footer_row(row):
            end_index = index
            break

        if _is_blank_row(row):
            consecutive_blank_rows += 1

            # A single blank row may occur inside an exported report.
            if consecutive_blank_rows >= 2:
                end_index = index - 1
                break
        else:
            consecutive_blank_rows = 0

    return _Section(
        worksheet=worksheet,
        title_row=title_index + 1,
        header_row=header_index + 1,
        data_start_row=header_index + 2,
        data_end_row=end_index + 1,
    )


def _is_positions_header(row: Sequence[Any]) -> bool:
    """Check whether a row resembles a MetaTrader Positions header."""
    headers = [_normalize_header(value) for value in row if _text(value)]

    counts: dict[str, int] = {}
    for header in headers:
        counts[header] = counts.get(header, 0) + 1

    has_time_pair = counts.get("time", 0) >= 2
    has_price_pair = counts.get("price", 0) >= 2

    required_groups = (
        {"position", "positionid", "ticket"},
        {"symbol", "instrument"},
        {"type", "direction", "side"},
        {"volume", "lots", "size"},
        {"profit", "pl", "pnl"},
    )

    has_required_groups = all(
        any(alias in counts for alias in aliases)
        for aliases in required_groups
    )

    return has_time_pair and has_price_pair and has_required_groups


# ---------------------------------------------------------------------------
# Column mapping
# ---------------------------------------------------------------------------

def _resolve_columns(header: Sequence[Any]) -> _Columns:
    """
    Resolve columns by their header names.

    For the standard MetaTrader Positions layout, the first Time and Price
    columns represent entry, and the second ones represent exit.
    """
    time_indices: list[int] = []
    price_indices: list[int] = []
    found: dict[str, int] = {}

    normalized_aliases = {
        key: {_normalize_header(alias) for alias in aliases}
        for key, aliases in HEADER_ALIASES.items()
    }

    for index, raw_value in enumerate(header):
        name = _normalize_header(raw_value)

        if not name:
            continue

        if name == "time":
            time_indices.append(index)
            continue

        if name == "price":
            price_indices.append(index)
            continue

        for logical_name, aliases in normalized_aliases.items():
            if name in aliases:
                if logical_name in found:
                    raise ImportValidationError(
                        f"Duplicate column for {logical_name!r} was found "
                        f"in the Positions header."
                    )
                found[logical_name] = index
                break

    missing: list[str] = []

    required = ("position", "symbol", "type", "volume", "profit")
    missing.extend(name for name in required if name not in found)

    if len(time_indices) < 2:
        missing.append("two Time columns (entry and exit)")

    if len(price_indices) < 2:
        missing.append("two Price columns (entry and exit)")

    if missing:
        raise ImportValidationError(
            "Missing required Positions columns: " + ", ".join(missing)
        )

    return _Columns(
        entry_time=time_indices[0],
        position_id=found["position"],
        symbol=found["symbol"],
        direction=found["type"],
        volume=found["volume"],
        entry_price=price_indices[0],
        exit_time=time_indices[1],
        exit_price=price_indices[1],
        stop_loss=found.get("sl"),
        take_profit=found.get("tp"),
        commission=found.get("commission"),
        swap=found.get("swap"),
        profit=found["profit"],
    )


# ---------------------------------------------------------------------------
# Cell parsing
# ---------------------------------------------------------------------------

def _get(row: Sequence[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _parse_datetime(value: Any, *, field: str, row_number: int) -> datetime:
    if isinstance(value, datetime):
        return value

    if value is None:
        raise ImportValidationError(
            f"Excel row {row_number}: {field} is empty."
        )

    value_text = _text(value)

    if not value_text:
        raise ImportValidationError(
            f"Excel row {row_number}: {field} is empty."
        )

    for date_format in DATETIME_FORMATS:
        try:
            return datetime.strptime(value_text, date_format)
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(value_text)
    except ValueError as exc:
        raise ImportValidationError(
            f"Excel row {row_number}: invalid {field} value "
            f"{value_text!r}."
        ) from exc


def _parse_decimal(
    value: Any,
    *,
    field: str,
    row_number: int,
    required: bool = True,
    default: Decimal | None = None,
) -> Decimal | None:
    if value is None or _text(value) == "":
        if required:
            raise ImportValidationError(
                f"Excel row {row_number}: {field} is empty."
            )
        return default

    if isinstance(value, bool):
        raise ImportValidationError(
            f"Excel row {row_number}: boolean is not a valid {field}."
        )

    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        text = _text(value).replace("\u2212", "-").replace("\u00a0", "")
        text = re.sub(r"\s+", "", text)

        # Handle common English-formatted thousands separators.
        if "," in text and "." in text:
            text = text.replace(",", "")
        elif re.fullmatch(r"-?\d{1,3}(,\d{3})+", text):
            text = text.replace(",", "")

        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ImportValidationError(
                f"Excel row {row_number}: invalid {field} value {value!r}."
            ) from exc

    if not number.is_finite():
        raise ImportValidationError(
            f"Excel row {row_number}: {field} must be a finite number."
        )

    return number


def _parse_trade(
    row: Sequence[Any],
    columns: _Columns,
    *,
    account_id: UUID,
    row_number: int,
) -> Trade:
    position_id = _text(_get(row, columns.position_id))
    symbol = _text(_get(row, columns.symbol))
    direction_text = _text(_get(row, columns.direction)).upper()

    if not position_id:
        raise ImportValidationError(
            f"Excel row {row_number}: Position identifier is empty."
        )

    if not symbol:
        raise ImportValidationError(
            f"Excel row {row_number}: Symbol is empty."
        )

    try:
        direction = TradeDirection(direction_text)
    except ValueError as exc:
        raise ImportValidationError(
            f"Excel row {row_number}: unsupported trade direction "
            f"{direction_text!r}; expected Buy or Sell."
        ) from exc

    entry_time = _parse_datetime(
        _get(row, columns.entry_time),
        field="entry_time",
        row_number=row_number,
    )

    entry_price = _parse_decimal(
        _get(row, columns.entry_price),
        field="entry_price",
        row_number=row_number,
    )

    volume = _parse_decimal(
        _get(row, columns.volume),
        field="volume",
        row_number=row_number,
    )

    exit_time_raw = _get(row, columns.exit_time)
    exit_price_raw = _get(row, columns.exit_price)

    exit_time = None
    if _text(exit_time_raw):
        exit_time = _parse_datetime(
            exit_time_raw,
            field="exit_time",
            row_number=row_number,
        )

    exit_price = None
    if _text(exit_price_raw):
        exit_price = _parse_decimal(
            exit_price_raw,
            field="exit_price",
            row_number=row_number,
        )

    if (exit_time is None) != (exit_price is None):
        raise ImportValidationError(
            f"Excel row {row_number}: exit_time and exit_price must either "
            f"both be present or both be empty."
        )

    stop_loss = _parse_decimal(
        _get(row, columns.stop_loss),
        field="stop_loss",
        row_number=row_number,
        required=False,
    )

    take_profit = _parse_decimal(
        _get(row, columns.take_profit),
        field="take_profit",
        row_number=row_number,
        required=False,
    )

    commission = _parse_decimal(
        _get(row, columns.commission),
        field="commission",
        row_number=row_number,
        required=False,
        default=Decimal("0"),
    )

    swap = _parse_decimal(
        _get(row, columns.swap),
        field="swap",
        row_number=row_number,
        required=False,
        default=Decimal("0"),
    )

    reported_profit = _parse_decimal(
        _get(row, columns.profit),
        field="reported_profit",
        row_number=row_number,
    )

    try:
        return Trade(
            account_id=account_id,
            symbol=symbol,
            direction=direction,
            volume=volume,
            entry_time=entry_time,
            entry_price=entry_price,
            position_id=position_id,
            stop_loss=stop_loss,
            take_profit=take_profit,
            exit_time=exit_time,
            exit_price=exit_price,
            commission=commission,
            swap=swap,
            reported_profit=reported_profit,
        )
    except (TypeError, ValueError) as exc:
        raise ImportValidationError(
            f"Excel row {row_number}: trade data violates the domain "
            f"model's validation rules: {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# Public importer
# ---------------------------------------------------------------------------

class ExcelImporter(TradeImporter):
    """Import a MetaTrader Positions table from an .xlsx report."""

    supported_extensions = frozenset(SUPPORTED_EXTENSIONS)

    def __init__(self, sheet_name: str | None = None) -> None:
        self._sheet_name = sheet_name

    def import_file(
        self,
        file_path: str | Path,
        account_id: UUID,
    ) -> ImportResult:
        if not isinstance(account_id, UUID):
            raise ImportValidationError(
                "account_id must be a UUID matching the domain model."
            )

        path = Path(file_path)

        if not path.exists():
            raise ImportValidationError(f"File not found: {path}")

        if not path.is_file():
            raise ImportValidationError(f"Path is not a file: {path}")

        if not self.supports(path):
            raise ImportValidationError(
                f"Unsupported Excel file extension {path.suffix!r}. "
                f"Supported extensions: {', '.join(sorted(self.supported_extensions))}."
            )

        try:
            workbook = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
            )
        except (InvalidFileException, OSError, ValueError) as exc:
            raise ImportValidationError(
                f"Could not open Excel workbook {path.name!r}: {exc}"
            ) from exc

        try:
            worksheet = self._select_worksheet(workbook)
            section = _find_positions_section(worksheet)

            header = tuple(
                worksheet.iter_rows(
                    min_row=section.header_row,
                    max_row=section.header_row,
                    values_only=True,
                )
            )[0]

            columns = _resolve_columns(header)

            trades: list[Trade] = []
            rows_read = 0

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=section.data_start_row,
                    max_row=section.data_end_row - 1,
                    values_only=True,
                ),
                start=section.data_start_row,
            ):
                if _is_blank_row(row):
                    continue

                if _is_footer_row(row):
                    break

                # Rows that contain only non-trade text are not treated as
                # transactions. A row that looks like a trade but has missing
                # required values will raise an explicit validation error.
                if not _looks_like_trade_row(row, columns):
                    if _is_non_trade_note(row):
                        continue
                    raise ImportValidationError(
                        f"Excel row {row_number}: non-empty row inside the "
                        f"Positions table could not be identified as a trade. "
                        f"Import cancelled to avoid silently losing data."
                    )

                trade = _parse_trade(
                    row,
                    columns,
                    account_id=account_id,
                    row_number=row_number,
                )
                trades.append(trade)
                rows_read += 1

            if not trades:
                raise ImportValidationError(
                    f"No trades were found in the Positions section of "
                    f"{worksheet.title!r}."
                )

            return ImportResult(
                account_id=account_id,
                trades=tuple(trades),
                source_path=str(path.resolve()),
                source_format="excel",
                rows_read=rows_read,
            )
        finally:
            workbook.close()

    def _select_worksheet(self, workbook: Any) -> Worksheet:
        if self._sheet_name is not None:
            if self._sheet_name not in workbook.sheetnames:
                raise ImportValidationError(
                    f"Worksheet {self._sheet_name!r} not found. "
                    f"Available worksheets: {workbook.sheetnames}"
                )

            worksheet = workbook[self._sheet_name]
            _find_positions_section(worksheet)
            return worksheet

        matching: list[Worksheet] = []

        for worksheet in workbook.worksheets:
            try:
                _find_positions_section(worksheet)
            except ImportValidationError:
                continue
            matching.append(worksheet)

        if not matching:
            raise ImportValidationError(
                "No worksheet containing a recognizable Positions section "
                f"was found. Available worksheets: {workbook.sheetnames}"
            )

        if len(matching) > 1:
            names = [worksheet.title for worksheet in matching]
            raise ImportValidationError(
                f"Multiple worksheets contain a Positions section: {names}. "
                f"Specify sheet_name explicitly to avoid importing the wrong data."
            )

        return matching[0]


def _looks_like_trade_row(row: Sequence[Any], columns: _Columns) -> bool:
    """Determine whether a row has enough identifying information to be a trade."""
    position_id = _text(_get(row, columns.position_id))
    symbol = _text(_get(row, columns.symbol))
    entry_time = _get(row, columns.entry_time)

    return bool(position_id or symbol or entry_time is not None)


def _is_non_trade_note(row: Sequence[Any]) -> bool:
    """Identify obvious report notes that are not trade rows."""
    non_empty = [_text(value) for value in row if _text(value)]

    if not non_empty:
        return True

    if len(non_empty) == 1:
        text = non_empty[0].strip().lower()
        return (
            text in SECTION_TERMINATORS
            or text.startswith("balance:")
            or text.startswith("equity:")
            or text.startswith("total:")
        )

    return False
